"""
从 Excel 模板导入内容到项目 JSON
运行: python3 scripts/import-content.py [excel_path]
默认读取: docs/content-template.xlsx
"""
import openpyxl
import json
import os
import sys

excel_path = sys.argv[1] if len(sys.argv) > 1 else 'docs/content-template.xlsx'

if not os.path.exists(excel_path):
    print(f'Error: {excel_path} not found')
    sys.exit(1)

wb = openpyxl.load_workbook(excel_path, data_only=True)

# === Import NEWS ===
if 'NEWS' in wb.sheetnames:
    ws = wb['NEWS']
    news = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        news.append({
            'id': str(row[0]),
            'title': str(row[1] or ''),
            'date': str(row[2] or ''),
            'summary': str(row[3] or ''),
            'content': str(row[4] or ''),
        })
    
    # Sort by id descending (newest first)
    news.sort(key=lambda x: int(x['id']), reverse=True)
    
    output = 'src/content/news.json'
    with open(output, 'w', encoding='utf-8') as f:
        json.dump(news, f, indent=2, ensure_ascii=False)
    print(f'✓ NEWS: {len(news)} articles -> {output}')

# === Import PRODUCTS ===
if 'PRODUCTS' in wb.sheetnames:
    ws = wb['PRODUCTS']
    
    # Group by category -> series
    categories = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0] or not row[1]:
            continue
        cat = str(row[0]).strip()
        series_name = str(row[1]).strip()
        image = str(row[2] or '').strip()
        spec_key = str(row[3] or '').strip()
        spec_value = str(row[4] or '').strip()
        
        if cat not in categories:
            categories[cat] = {}
        if series_name not in categories[cat]:
            categories[cat][series_name] = {'image': image, 'specs': {}}
        if spec_key and spec_value:
            categories[cat][series_name]['specs'][spec_key] = spec_value
    
    # Merge into existing JSON files
    products_dir = 'src/content/products'
    for cat_slug, series_dict in categories.items():
        json_path = os.path.join(products_dir, f'{cat_slug}.json')
        
        if os.path.exists(json_path):
            with open(json_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
        else:
            data = {'name': cat_slug, 'slug': cat_slug, 'parent': None, 'series': []}
        
        # Update or add series
        existing_names = {s['name'] for s in data['series']}
        for name, info in series_dict.items():
            slug = name.lower().strip()
            slug = slug.replace(' ', '-').replace('/', '-').replace('(', '').replace(')', '')
            
            if name in existing_names:
                # Update existing
                for s in data['series']:
                    if s['name'] == name:
                        if info['image']:
                            s['image'] = info['image']
                        s['specs'].update(info['specs'])
                        break
            else:
                # Add new
                data['series'].append({
                    'name': name,
                    'slug': slug,
                    'image': info['image'],
                    'images': [info['image']] if info['image'] else [],
                    'specs': info['specs']
                })
        
        with open(json_path, 'w', encoding='utf-8') as f:
            json.dump(data, f, indent=2, ensure_ascii=False)
        print(f'✓ PRODUCTS/{cat_slug}: {len(series_dict)} series updated -> {json_path}')

wb.close()
print('\nDone! Run "npm run build" to rebuild the site.')
