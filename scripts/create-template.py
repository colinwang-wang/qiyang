"""
生成内容更新 Excel 模板
运行: python3 scripts/create-template.py
输出: docs/content-template.xlsx
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

wb = openpyxl.Workbook()

# === NEWS Sheet ===
ws_news = wb.active
ws_news.title = 'NEWS'

headers = ['id', 'title', 'date', 'summary', 'content']
header_fill = PatternFill(start_color='1E40AF', end_color='1E40AF', fill_type='solid')
header_font = Font(color='FFFFFF', bold=True)

for col, h in enumerate(headers, 1):
    cell = ws_news.cell(row=1, column=col, value=h)
    cell.fill = header_fill
    cell.font = header_font

# Example row
ws_news.cell(row=2, column=1, value='1')
ws_news.cell(row=2, column=2, value='Welcome to Qiyang Electronics')
ws_news.cell(row=2, column=3, value='2026-05-07')
ws_news.cell(row=2, column=4, value='Our official website is now online.')
ws_news.cell(row=2, column=5, value='Full article content here...')

ws_news.column_dimensions['A'].width = 5
ws_news.column_dimensions['B'].width = 35
ws_news.column_dimensions['C'].width = 12
ws_news.column_dimensions['D'].width = 45
ws_news.column_dimensions['E'].width = 60

# === PRODUCTS Sheet ===
ws_prod = wb.create_sheet('PRODUCTS')

prod_headers = ['category', 'series_name', 'image_filename', 'spec_key', 'spec_value']
for col, h in enumerate(prod_headers, 1):
    cell = ws_prod.cell(row=1, column=col, value=h)
    cell.fill = header_fill
    cell.font = header_font

# Example rows
examples = [
    ('rectangular-connector', 'J63A series', 'image1.png', 'Operating Temperature', '-55℃ ~ +125℃'),
    ('rectangular-connector', 'J63A series', 'image1.png', 'Rated Current', '1A'),
    ('rectangular-connector', 'J63A series', 'image1.png', 'Mechanical Life', '500 cycles'),
    ('rectangular-connector', 'J30 series', 'image3.png', 'Operating Temperature', '-55℃ ~ +185℃'),
]
for i, row in enumerate(examples, 2):
    for col, val in enumerate(row, 1):
        ws_prod.cell(row=i, column=col, value=val)

ws_prod.column_dimensions['A'].width = 25
ws_prod.column_dimensions['B'].width = 25
ws_prod.column_dimensions['C'].width = 15
ws_prod.column_dimensions['D'].width = 30
ws_prod.column_dimensions['E'].width = 30

# === 说明 Sheet ===
ws_help = wb.create_sheet('说明')
instructions = [
    ['内容更新 Excel 模板使用说明'],
    [''],
    ['NEWS 表:'],
    ['  - 每行一篇文章，id 递增，最新的放最后'],
    ['  - date 格式: YYYY-MM-DD'],
    ['  - 填完后运行: python3 scripts/import-content.py'],
    [''],
    ['PRODUCTS 表:'],
    ['  - category: 产品类别 slug (如 rectangular-connector, rf-connector 等)'],
    ['  - series_name: 系列名称'],
    ['  - image_filename: 图片文件名 (放到 public/images/products/ 下)'],
    ['  - spec_key / spec_value: 规格参数，同一系列可多行'],
    [''],
    ['可用的 category 值:'],
    ['  rectangular-connector, circular-connector, rf-connector,'],
    ['  cable-assembly, bldc-torque-motor, resolver-transmitter, rotary-encoder'],
]
for i, row in enumerate(instructions, 1):
    ws_help.cell(row=i, column=1, value=row[0] if row else '')
ws_help.column_dimensions['A'].width = 70

wb.save('docs/content-template.xlsx')
print('Created: docs/content-template.xlsx')
